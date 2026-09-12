"""Etiket listesi (Excel) → cihaz ürün kodu.

Belgenin yanında duran `* Etiket *.xlsx` dosyası panodaki her cihaz etiketi için
malzeme kodlarını taşır: `Etiket | Kod | Adet | Ort`. Bir etikette birden çok kod olur;
sigortada üç sigorta buşonu (adet 3) ile kutusu (adet 1) aynı etikette durur.

Hangi kodun CİHAZIN KENDİSİ olduğu seçilebilir olmalı, çünkü müşteriye göre değişir:
seçim kuralı müşteri profilinde durur. Varsayılan `first_qty_1` — ölçüm (Troester 13SB003,
285 etiket): sigortada RIT kutusunu, rölede ana cihazı, motor korumada şalterin kendisini
seçiyor; yardımcı bloklar ve buşonlar `other_codes` altında durur, atılmaz.
"""
from collections import defaultdict
import re

TAG = re.compile(r'^=?[A-Za-z0-9]+\+?[A-Za-z0-9]*-[A-Za-z0-9._]+$')


def find_file(folder):
    """Klasördeki etiket dosyası. Birden çoksa en yenisi; yoksa None."""
    from pathlib import Path
    folder = Path(folder)
    if not folder.exists():
        return None
    files = [p for p in folder.glob('*.xlsx') if 'etiket' in p.name.lower()]
    if not files:
        return None
    return max(files, key=lambda p: p.stat().st_mtime)


def read(path):
    """`{etiket: [{code, qty, location, sheet, row}]}`. Başlık satırı atlanır."""
    import openpyxl
    book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    out = defaultdict(list)
    for sheet_name in book.sheetnames:
        sheet = book[sheet_name]
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = ['' if v is None else str(v).strip() for v in row]
            if len(cells) < 2 or not cells[0] or not cells[1]:
                continue
            if not TAG.match(cells[0]):
                continue                       # başlık ve boş satırlar
            quantity = None
            if len(cells) > 2 and cells[2]:
                try:
                    quantity = int(float(cells[2].replace(',', '.')))
                except ValueError:
                    quantity = None
            out[cells[0]].append(dict(code=cells[1], qty=quantity,
                                      location=cells[3] if len(cells) > 3 else '',
                                      sheet=sheet_name, row=index))
    book.close()
    return dict(out)


def pick(rows, strategy='first_qty_1', prefer_prefix=None):
    """Cihazın kendi kodu. Karar açıklanabilir olsun diye gerekçe de döner.

    `first_qty_1` — adedi 1 olan ilk kod. Çoklu parçalı etikette (sigorta buşonu x3 +
    kutu x1) kutuyu seçer; yardımcı blok gibi ikinci kodlar listede kalır.
    `prefer_prefix` — verilirse önce bu üretici ön ekinden adedi 1 olan kod aranır.
    """
    if not rows:
        return None
    if prefer_prefix:
        wanted = [r for r in rows if r['code'].upper().startswith(prefer_prefix.upper())
                  and r['qty'] == 1]
        if wanted:
            return dict(wanted[0], reason='ön ek %s ve adet 1' % prefer_prefix)
    if strategy == 'first_row':
        return dict(rows[0], reason='listedeki ilk kod')
    ones = [r for r in rows if r['qty'] == 1]
    if ones:
        return dict(ones[0], reason='adedi 1 olan ilk kod')
    return dict(rows[0], reason='adedi 1 olan kod yok; listedeki ilk kod')


def for_device(table, device_tag, strategy='first_qty_1', prefer_prefix=None):
    """Bir cihazın kodu. Etiket listede yoksa None — uydurulmaz."""
    rows = table.get(device_tag)
    if not rows:
        return None
    chosen = pick(rows, strategy, prefer_prefix)
    return dict(code=chosen['code'], qty=chosen['qty'], reason=chosen['reason'],
                source='LABEL_LIST', sheet=chosen['sheet'],
                other_codes=[dict(code=r['code'], qty=r['qty']) for r in rows
                             if r['code'] != chosen['code']])
